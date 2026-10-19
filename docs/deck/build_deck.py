# /// script
# requires-python = ">=3.10"
# dependencies = ["python-pptx==1.0.2"]
# ///
"""Build the Arabic AI Video Search use-case deck.

Run from the repository root:

    uv run docs/deck/build_deck.py

The deck is written next to this script.
Every fact and number comes from the PoC technical plan (its report and board),
and the demo slide describes the app as the project README does;
the speaker notes name the plan's sections and the upstream links.
Illustrations are drawn from shapes: no photos of real people are used.
"""

from __future__ import annotations

import datetime as dt
import re
from dataclasses import dataclass, field
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

OUT = Path(__file__).with_name("arabic-video-search-use-cases.pptx")

# Fonts: Aptos for Latin text and Dubai for Arabic.
# Both ship with Microsoft 365 on Windows and macOS, and Dubai shapes Arabic correctly.
LATIN_FONT = "Aptos"
ARABIC_FONT = "Dubai"

NAVY = "0F2A44"
INK = "1F2937"
MUTED = "5B6573"
TEAL = "0E7C7B"
TEAL_LIGHT = "E3F2F1"
AMBER = "B45F06"
AMBER_LIGHT = "FDF1E1"
PANEL = "F3F5F7"
LINE = "D5DBE1"
WHITE = "FFFFFF"
FRAME = "2B3A4A"
FRAME_FLOOR = "3A4B5C"
SHIRT_RED = "C0392B"
MAROON = "7B2D3B"
BAG_BLACK = "111111"
JACKET_BLUE = "2E6FD8"
TROUSER = "4A5563"
SKIN = "C9A27E"
GREY_TOP = "8A96A3"
GREEN_TOP = "3F8F5B"

SOURCE_LINE = (
    "Source: the Arabic AI Video Search PoC technical plan (its report and board); "
    "the decisions D1 to D7 are summarised in docs/decisions.md."
)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN = 0.6

ARABIC_RE = re.compile(r"[؀-ۿݐ-ݿﭐ-﷿ﹰ-﻿]")


def rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_)


# ---------------------------------------------------------------- text


@dataclass
class R:
    """A run of text with optional overrides of its paragraph's style."""

    text: str
    bold: bool | None = None
    color: str | None = None
    size: float | None = None
    italic: bool | None = None


@dataclass
class P:
    """A paragraph: plain text or a list of runs."""

    runs: str | list[R]
    size: float = 16
    color: str = INK
    bold: bool = False
    align: str = "auto"  # auto: right for right-to-left paragraphs, else left
    rtl: bool | None = None  # None: right-to-left when the text is mostly Arabic
    space_after: float = 4
    space_before: float = 0
    line_spacing: float | None = None
    italic: bool = False

    def run_list(self) -> list[R]:
        return [R(self.runs)] if isinstance(self.runs, str) else self.runs

    def is_rtl(self) -> bool:
        if self.rtl is not None:
            return self.rtl
        text = "".join(r.text for r in self.run_list())
        letters = [c for c in text if c.isalpha()]
        arabic = sum(1 for c in letters if ARABIC_RE.match(c))
        return bool(letters) and arabic / len(letters) > 0.5


ARABIC_SPAN_RE = re.compile(
    r"[\u0600-\u06FF\u0750-\u077F\uFB50-\uFDFF\uFE70-\uFEFF]+"
    r"(?:\s+[\u0600-\u06FF\u0750-\u077F\uFB50-\uFDFF\uFE70-\uFEFF]+)*"
)
LATIN_SPAN_RE = re.compile(r"[A-Za-z0-9]+(?:[\s\-.:'\"]+[A-Za-z0-9]+)*")


def split_by_script(text_: str, rtl: bool) -> list[str]:
    """Split a run into Arabic and Latin pieces.

    Each piece gets its own language tag, so PowerPoint orders mixed lines correctly.
    Neutral characters such as spaces and dots stay with the paragraph's own direction.
    """
    pattern = LATIN_SPAN_RE if rtl else ARABIC_SPAN_RE
    pieces, pos = [], 0
    for m in pattern.finditer(text_):
        if m.start() > pos:
            pieces.append(text_[pos : m.start()])
        pieces.append(m.group())
        pos = m.end()
    if pos < len(text_):
        pieces.append(text_[pos:])
    return pieces or [text_]


def style_run(run, *, size, color, bold, italic, text):
    font = run.font
    font.size = Pt(size)
    font.bold = bold
    font.italic = italic
    font.color.rgb = rgb(color)
    font.name = LATIN_FONT
    rpr = run._r.get_or_add_rPr()
    rpr.set("lang", "ar-JO" if ARABIC_RE.search(text) else "en-GB")
    latin = rpr.find(qn("a:latin"))
    for tag, face in (("a:cs", ARABIC_FONT), ("a:ea", LATIN_FONT)):
        el = rpr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rpr, qn(tag))
            latin.addnext(el)
        el.set("typeface", face)


def fill_text_frame(tf, paras: list[P], *, anchor="top", margins=(0, 0, 0, 0)):
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = (
        Inches(m) for m in margins
    )
    tf.vertical_anchor = {
        "top": MSO_ANCHOR.TOP,
        "middle": MSO_ANCHOR.MIDDLE,
        "bottom": MSO_ANCHOR.BOTTOM,
    }[anchor]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        rtl = p.is_rtl()
        align = p.align if p.align != "auto" else ("r" if rtl else "l")
        para.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        if rtl:
            para._p.get_or_add_pPr().set("rtl", "1")
        para.space_after = Pt(p.space_after)
        para.space_before = Pt(p.space_before)
        if p.line_spacing:
            para.line_spacing = p.line_spacing
        for r in p.run_list():
            for piece in split_by_script(r.text, rtl):
                run = para.add_run()
                run.text = piece
                style_run(
                    run,
                    size=r.size or p.size,
                    color=r.color or p.color,
                    bold=p.bold if r.bold is None else r.bold,
                    italic=p.italic if r.italic is None else r.italic,
                    text=piece,
                )


def text(slide, x, y, w, h, paras, *, anchor="top", margins=(0, 0, 0, 0)):
    if isinstance(paras, P):
        paras = [paras]
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    fill_text_frame(shape.text_frame, paras, anchor=anchor, margins=margins)
    return shape


# ---------------------------------------------------------------- shapes


def shape(slide, kind, x, y, w, h, *, fill=None, line=None, line_w=1.0, radius=None, dash=None):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(line_w)
        if dash:
            s.line.dash_style = dash
    else:
        s.line.fill.background()
    if radius is not None and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    return s


def box(slide, x, y, w, h, *, fill=PANEL, line=None, radius=0.08, paras=None, anchor="top", pad=0.18, line_w=1.0):
    s = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=fill, line=line, radius=radius, line_w=line_w)
    if paras is not None:
        fill_text_frame(
            s.text_frame,
            paras if isinstance(paras, list) else [paras],
            anchor=anchor,
            margins=(pad, pad * 0.8, pad, pad * 0.8),
        )
    return s


def circle_label(slide, x, y, d, label, *, fill=TEAL, color=WHITE, size=14):
    s = shape(slide, MSO_SHAPE.OVAL, x, y, d, d, fill=fill)
    fill_text_frame(s.text_frame, [P(label, size=size, color=color, bold=True, align="c", space_after=0)], anchor="middle")
    return s


def connector(slide, x1, y1, x2, y2):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    style = c._element.find(qn("p:style"))
    if style is not None:
        c._element.remove(style)  # the default style adds a theme shadow
    return c


def plain_line(slide, x1, y1, x2, y2, *, color=LINE, width=1.0, dash=None):
    c = connector(slide, x1, y1, x2, y2)
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    if dash:
        c.line.dash_style = dash
    return c


def chevron(slide, x, y, w=0.32, h=0.42, color=LINE):
    return shape(slide, MSO_SHAPE.CHEVRON, x, y, w, h, fill=color)


def person(slide, x, y, h, *, top=GREY_TOP, bottom=TROUSER, bag=None, bag_side="r", head=SKIN, outline=None):
    """A simple standing figure, x/y is its top-left corner, h its height (inches)."""
    w = 0.42 * h
    cx = x + w / 2
    hd = 0.19 * h
    shape(slide, MSO_SHAPE.OVAL, cx - hd / 2, y, hd, hd, fill=head, line=outline, line_w=0.75)
    tw, th = 0.40 * h, 0.40 * h
    torso = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, cx - tw / 2, y + 0.21 * h, tw, th, fill=top, radius=0.3, line=outline, line_w=0.75)
    lw = 0.30 * h
    shape(slide, MSO_SHAPE.RECTANGLE, cx - lw / 2, y + 0.58 * h, lw, 0.42 * h, fill=bottom, line=outline, line_w=0.75)
    shape(slide, MSO_SHAPE.RECTANGLE, cx - 0.012 * h, y + 0.68 * h, 0.024 * h, 0.32 * h, fill=FRAME if outline else "222B35")
    if bag:
        bw, bh = 0.17 * h, 0.2 * h
        bx = cx + tw / 2 - 0.02 * h if bag_side == "r" else cx - tw / 2 - bw + 0.02 * h
        shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, bx, y + 0.5 * h, bw, bh, fill=bag, radius=0.2)
    return torso


def suitcase(slide, x, y, h, color=BAG_BLACK):
    w = 0.62 * h
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x + w * 0.32, y, w * 0.36, h * 0.18, fill=None, line=color, line_w=1.5, radius=0.3)
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y + h * 0.14, w, h * 0.86, fill=color, radius=0.12)


def camera_icon(slide, x, y, s, color=NAVY):
    shape(slide, MSO_SHAPE.RECTANGLE, x + 0.08 * s, y, 0.1 * s, 0.32 * s, fill=color)
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y + 0.26 * s, 0.78 * s, 0.38 * s, fill=color, radius=0.25)
    shape(slide, MSO_SHAPE.TRAPEZOID, x + 0.74 * s, y + 0.28 * s, 0.26 * s, 0.34 * s, fill=color).rotation = 90


def frame(slide, x, y, w, h, *, label=None, people=(), floor=True):
    """A CCTV-style video frame with simple figures inside.

    people: list of (rel_x, rel_h, kwargs) with rel_x as a fraction of the width
    and rel_h as the figure height relative to the frame height.
    """
    shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, h, fill=FRAME)
    if floor:
        shape(slide, MSO_SHAPE.RECTANGLE, x, y + h * 0.62, w, h * 0.38, fill=FRAME_FLOOR)
    for rel_x, rel_h, kw in people:
        ph = h * rel_h
        if kw.pop("_suitcase", False):
            suitcase(slide, x + w * rel_x, y + h * 0.9 - ph, ph, kw.get("color", BAG_BLACK))
        else:
            person(slide, x + w * rel_x, y + h * 0.92 - ph, ph, outline=None, **kw)
    if label:
        text(
            slide, x + 0.06, y + 0.04, w - 0.12, 0.26,
            P(label, size=9, color=WHITE, bold=True, space_after=0),
        )


# ---------------------------------------------------------------- slide chrome


@dataclass
class Deck:
    prs: Presentation
    total: int
    count: int = 0
    slides: list = field(default_factory=list)

    def new(self, *, kicker=None, title=None, arabic=None, notes=None, dark=False):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        if not dark:
            self.count += 1
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = rgb(NAVY if dark else WHITE)
        if not dark:
            shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, 13.333, 0.12, fill=TEAL)
            plain_line(slide, MARGIN, 7.02, 13.333 - MARGIN, 7.02, color=LINE)
            text(
                slide, MARGIN, 7.08, 8, 0.3,
                P("Arabic AI Video Search  ·  PoC use cases", size=10, color=MUTED, space_after=0),
            )
            text(
                slide, 13.333 - MARGIN - 2, 7.08, 2, 0.3,
                P(f"{self.count} / {self.total}", size=10, color=MUTED, align="r", space_after=0),
            )
        if kicker:
            text(slide, MARGIN, 0.38, 8.4, 0.32, P(kicker.upper(), size=12, color=TEAL, bold=True, space_after=0))
        if arabic:
            text(
                slide, 13.333 - MARGIN - 4.2, 0.3, 4.2, 0.45,
                P(arabic, size=18, color=TEAL, bold=True, space_after=0, rtl=True),
            )
        if title:
            text(slide, MARGIN, 0.72, 12.1, 1.0, P(title, size=28, color=NAVY, bold=True, space_after=0, line_spacing=0.95))
        if notes:
            slide.notes_slide.notes_text_frame.text = notes.strip() + "\n\n" + SOURCE_LINE
        self.slides.append(slide)
        return slide


def numbered_steps(slide, x, y, w, steps, *, gap=0.16, size=15, lead_size=16, num_d=0.42, heights=None):
    """Vertical list of numbered steps. steps: list of (lead, body)."""
    cy = y
    for i, (lead, body) in enumerate(steps):
        h = heights[i] if heights else 0.9
        circle_label(slide, x, cy + 0.02, num_d, str(i + 1), size=14)
        paras = [P(lead, size=lead_size, bold=True, color=NAVY, space_after=2)]
        if body:
            paras += body if isinstance(body, list) else [P(body, size=size, color=INK, space_after=0)]
        text(slide, x + num_d + 0.18, cy, w - num_d - 0.18, h, paras)
        cy += h + gap
    return cy


# ---------------------------------------------------------------- slides


def slide_title(deck: Deck):
    s = deck.new(
        dark=True,
        notes="""
Opening.
This deck shows how security teams would use Arabic video search, as three search stories: a suspect, a left bag, and a lost child.
The technical choices behind it are settled in the PoC technical plan (decisions D1 to D7).
The query on the right is the plan's headline example: "a man wearing a red shirt".
""",
    )
    shape(s, MSO_SHAPE.RECTANGLE, 0, 0, 0.18, 7.5, fill=TEAL)
    text(s, 0.9, 1.2, 7.4, 0.4, P("PROOF OF CONCEPT  ·  USE CASES", size=14, color="7FD1CC", bold=True, space_after=0))
    text(s, 0.9, 1.7, 7.6, 1.1, P("Arabic AI Video Search", size=48, color=WHITE, bold=True, space_after=0))
    text(
        s, 0.9, 2.75, 7.4, 0.7,
        P("البحث الذكي في تسجيلات كاميرات المراقبة باللغة العربية", size=24, color="BFE6E3", rtl=True, align="l", space_after=0),
    )
    text(
        s, 0.9, 3.75, 7.0, 1.4,
        [
            P("Search CCTV recordings by describing a person in Arabic.", size=20, color=WHITE, space_after=6),
            P("How a security team would use it: a suspect, a left bag and a lost child.", size=20, color="C9D3DD", space_after=0),
        ],
    )
    text(s, 0.9, 6.45, 7, 0.4, P("October 2026  ·  Discussion draft", size=13, color="9FB0C0", space_after=0))

    # Search bar and two result frames, as a simple illustration.
    x0, y0, w0 = 8.45, 1.55, 4.25
    box(
        s, x0, y0, w0, 0.95, fill=WHITE, radius=0.18,
        paras=[P("رجل يرتدي قميص أحمر", size=24, color=NAVY, bold=True, rtl=True, space_after=0)],
        anchor="middle", pad=0.25,
    )
    text(s, x0, y0 + 1.02, w0, 0.32, P("\"a man wearing a red shirt\"", size=13, color="C9D3DD", italic=True, align="r", space_after=0))
    fy = y0 + 1.6
    for i, (cam, t) in enumerate((("CAM-03  14:02:17", "Main entrance"), ("CAM-07  14:05:40", "Food court"))):
        fx = x0 + i * 2.2
        frame(s, fx, fy, 2.05, 1.45, label=cam, people=[(0.12, 0.62, dict(top=GREY_TOP)), (0.55, 0.78, dict(top=SHIRT_RED, bag=BAG_BLACK))])
        text(s, fx, fy + 1.52, 2.05, 0.3, P(t, size=12, color="C9D3DD", space_after=0))
    box(
        s, x0, fy + 2.0, w0, 0.6, fill="153A5C", radius=0.2,
        paras=[P([R("A person confirms every match", bold=True, color=WHITE)], size=14, align="c", space_after=0)],
        anchor="middle",
    )


def slide_problem(deck: Deck):
    s = deck.new(
        kicker="The problem",
        arabic="المشكلة",
        title="After an incident, the answer sits in hours of footage from many cameras, searched by hand",
        notes="""
The point: the search is manual and grows with every camera and every hour.
The 100 hours figure is our arithmetic from the plan's own example: a missing-child case covering 50 nearby cameras for 2 hours (report section 4, "City scale").
English-only: NVIDIA's open Video Search and Summarization blueprint uses the same search shape but works in English only (report section 1; https://docs.nvidia.com/vss/3.2.1/agent-workflow-search.html).
No public benchmark covers Jordanian Arabic, and no Arabic text-to-person dataset exists (report section 2).
""",
    )
    # Wall of camera tiles.
    gx, gy, tw, th, gap = 0.6, 2.05, 1.38, 0.95, 0.12
    people_sets = [
        [(0.2, 0.6, dict(top=GREY_TOP))],
        [(0.5, 0.7, dict(top=GREEN_TOP))],
        [],
        [(0.15, 0.55, dict(top=MAROON)), (0.6, 0.62, dict(top=GREY_TOP))],
        [(0.4, 0.66, dict(top=JACKET_BLUE))],
        [(0.65, 0.5, dict(top=GREY_TOP))],
        [(0.3, 0.7, dict(top=SHIRT_RED, bag=BAG_BLACK))],
        [],
        [(0.1, 0.6, dict(top=GREEN_TOP)), (0.55, 0.5, dict(top=GREY_TOP))],
        [(0.45, 0.68, dict(top=GREY_TOP))],
        [],
        [(0.25, 0.62, dict(top=JACKET_BLUE))],
    ]
    for i in range(12):
        r, c = divmod(i, 4)
        frame(s, gx + c * (tw + gap), gy + r * (th + gap), tw, th, label=f"CAM-{i + 1:02d}", people=people_sets[i])
    text(
        s, gx, gy + 3 * (th + gap) + 0.05, 4 * tw + 3 * gap, 0.6,
        P("Each camera's recording is scrubbed in turn, for the whole time window.", size=13, color=MUTED, italic=True, space_after=0),
    )

    x = 6.85
    w = 5.85
    points = [
        ("Many cameras, one description", "A witness or a parent describes a person. Staff then play back camera after camera, looking for that person."),
        ("The work grows with every camera and hour", [
            P([R("A lost-child case covering "), R("50 nearby cameras for 2 hours", bold=True), R(" is "), R("100 hours of video", bold=True, color=TEAL), R(" to watch.")], size=15, space_after=0),
        ]),
        ("Search tools that understand descriptions work in English", "The closest open design, NVIDIA's video-search blueprint, works in English only, and no public test covers Jordanian Arabic."),
    ]
    numbered_steps(s, x, 2.05, w, points, heights=[1.05, 1.05, 1.05], gap=0.3)


def slide_what(deck: Deck):
    s = deck.new(
        kicker="What the system does",
        arabic="ماذا يفعل النظام",
        title="Type a description in Arabic. Get the matching clips, with camera, place and time.",
        notes="""
One picture of the whole system, left to right.
1. The officer types a description in Arabic, formal or dialect. Example from the plan: "رجل يرتدي قميص أحمر", a man wearing a red shirt.
2. A local language model turns it into clear fields (person, clothing colours, bag) plus a clean English description (decision D1, hybrid; D4, local only).
3. The search is narrowed to the chosen cameras and time window, then matched against a catalogue of every person already seen in the footage.
4. Results show a short clip, camera, place, time, and why it matched.
5. A person confirms or rejects every result: a result is a lead, not proof (report section 5).
How it works underneath (report section 1): indexing decodes the footage, detects and tracks people, keeps a few best crops per person, describes each one, and stores it with camera, place, time and a pointer to the clip.
""",
    )
    y = 2.0
    cols = [0.6, 3.15, 5.7, 8.25, 10.8]
    cw = 1.95
    heads = [
        ("1", "Type", "اكتب الوصف"),
        ("2", "Understand", "فهم الوصف"),
        ("3", "Search", "البحث"),
        ("4", "Results", "النتائج"),
        ("5", "Confirm", "التأكيد"),
    ]
    for i, (n, en, ar) in enumerate(heads):
        x = cols[i]
        circle_label(s, x, y, 0.46, n, fill=AMBER if i == 4 else TEAL)
        text(s, x + 0.56, y - 0.02, cw, 0.32, P(en, size=17, bold=True, color=NAVY, space_after=0))
        text(s, x + 0.56, y + 0.27, cw - 0.56, 0.3, P(ar, size=12, color=MUTED, rtl=True, align="l", space_after=0))
        if i < 4:
            chevron(s, cols[i + 1] - 0.42, y + 1.55, 0.26, 0.42)

    py = 2.75
    # 1. Type
    box(
        s, cols[0], py, cw + 0.15, 1.1, fill=WHITE, line=TEAL, line_w=1.75, radius=0.15,
        paras=[P("رجل يرتدي قميص أحمر", size=14, bold=True, color=NAVY, rtl=True, align="c", space_after=0)],
        anchor="middle", pad=0.08,
    )
    text(s, cols[0], py + 1.18, cw + 0.15, 0.6, P("\"a man wearing a red shirt\", in formal Arabic or dialect", size=12, color=MUTED, italic=True, space_after=0))
    # 2. Understand
    for j, (ar, en) in enumerate((("رجل", "man"), ("قميص أحمر", "red shirt"))):
        box(
            s, cols[1], py + j * 0.62, cw + 0.15, 0.5, fill=TEAL_LIGHT, radius=0.5,
            paras=[P([R(ar, bold=True, color=NAVY), R("  ·  "), R(en)], size=14, align="c", space_after=0, rtl=False)],
            anchor="middle", pad=0.08,
        )
    text(s, cols[1], py + 1.32, cw + 0.15, 1.2, P("A local language model turns the words, including dialect, into clear fields and plain English.", size=12, color=MUTED, space_after=0))
    # 3. Search
    for j in range(3):
        camera_icon(s, cols[2] + 0.05 + j * 0.68, py + 0.12, 0.55, color=NAVY if j < 2 else LINE)
    box(s, cols[2], py + 0.75, cw + 0.15, 0.48, fill=PANEL, radius=0.5,
        paras=[P("8 cameras · 13:00-15:00", size=12, color=NAVY, align="c", space_after=0)], anchor="middle", pad=0.06)
    text(s, cols[2], py + 1.32, cw + 0.15, 1.2, P("Only the cameras and hours that matter are searched, in a catalogue of every person already seen.", size=12, color=MUTED, space_after=0))
    # 4. Results
    for j, (cam, place, top) in enumerate((("CAM-03", "Entrance · 14:02", SHIRT_RED), ("CAM-07", "Food court · 14:05", SHIRT_RED))):
        frame(s, cols[3], py + j * 0.7, 0.85, 0.6, people=[(0.3, 0.8, dict(top=top))])
        text(s, cols[3] + 0.93, py + j * 0.7 + 0.02, 1.3, 0.6,
             [P(cam, size=12, bold=True, color=NAVY, space_after=0), P(place, size=11, color=MUTED, space_after=0)])
    text(s, cols[3], py + 1.47, cw + 0.15, 1.2, P("Short clips, ranked, each with camera, place, time and why it matched.", size=12, color=MUTED, space_after=0))
    # 5. Confirm
    box(s, cols[4], py, cw, 0.5, fill=TEAL_LIGHT, radius=0.5,
        paras=[P("✓  Confirm", size=14, bold=True, color=TEAL, align="c", space_after=0)], anchor="middle", pad=0.06)
    box(s, cols[4], py + 0.62, cw, 0.5, fill=AMBER_LIGHT, radius=0.5,
        paras=[P("✕  Reject", size=14, bold=True, color=AMBER, align="c", space_after=0)], anchor="middle", pad=0.06)
    text(s, cols[4], py + 1.32, cw, 1.2, P("A person decides. Every result is a lead, not proof, and every search is logged.", size=12, color=MUTED, space_after=0))

    box(
        s, 0.6, 5.35, 12.13, 0.85, fill=PANEL, radius=0.1,
        paras=[P([
            R("Underneath: ", bold=True, color=NAVY),
            R("each recording is first turned into a catalogue of every person seen, with clothing, colours and bags, camera, place and time. "
              "A query searches that catalogue and jumps to the right second of video."),
        ], size=14, space_after=0)],
        anchor="middle", pad=0.25,
    )


def slide_stories(deck: Deck):
    s = deck.new(
        kicker="Three stories",
        arabic="ثلاث حالات استخدام",
        title="One search engine, three ways a security team would use it",
        notes="""
All three stories are appearance searches with the same engine; they differ in what the target looks like and how hard it is to recognise (decision D6, all three as search stories).
Where they apply, from the original brief: finding missing persons or children, assisting investigations, and security in airports, tourism areas and government buildings.
Suspect: the core search, no extra work. The query is the plan's Jordanian-dialect example, "a guy wearing a red top, with a black bag"; the query model must read dialect colour forms such as حمرا and سودا (report section 2).
Left bag: a search the officer starts, small extra work, since backpack, handbag and suitcase are classes the detector already knows. Small bags far from the camera get missed. An automatic left-bag alarm is a separate feature and is not part of the PoC (report D6 answer).
Lost child: medium extra work, because age is hard to read from a body, so the system shows candidates and an officer reviews each one. It searches clothing, colours and size, never faces (decision D5). The query is the plan's example, "a small child wearing a blue jacket".
The pictures are illustrations, drawn from shapes.
""",
    )
    cards = [
        ("1", "Finding a suspect", "البحث عن مشتبه به",
         "A witness saw a man in a red top with a black bag. Where is he on the site's cameras?",
         "زلمة لابس بلوزة حمرا ومعه شنطة سودا", "\"a guy in a red top, with a black bag\" (Jordanian)",
         "Ranked clips of him with camera, place and time. An officer confirms one, then follows him camera to camera.",
         [(0.36, 0.86, dict(top=SHIRT_RED, bag=BAG_BLACK))]),
        ("2", "Finding who left a bag", "من ترك الحقيبة؟",
         "An unattended suitcase is found in a lobby. Who brought it, and where did they go?",
         "حقيبة سفر سوداء", "\"a black suitcase\"",
         "The moment the bag arrived, and the person who brought it. Their clothing becomes the next search.",
         [(0.2, 0.84, dict(top=GREY_TOP)), (0.62, 0.34, dict(_suitcase=True))]),
        ("3", "Finding a lost child", "البحث عن طفل مفقود",
         "A parent has lost a small child in a blue jacket near Gate 3.",
         "طفل صغير لابس جاكيت أزرق", "\"a small child wearing a blue jacket\"",
         "The latest sightings near Gate 3, newest first, as candidates an officer checks. Clothing and size, never faces.",
         [(0.18, 0.86, dict(top=GREY_TOP)), (0.58, 0.52, dict(top=JACKET_BLUE))]),
    ]
    cw, gap, y, h = 3.85, 0.29, 1.9, 4.95
    for i, (n, en, ar, situation, q, gloss, result, ppl) in enumerate(cards):
        x = 0.6 + i * (cw + gap)
        ix, iw = x + 0.22, cw - 0.44
        box(s, x, y, cw, h, fill=PANEL, radius=0.04)
        frame(s, x, y, cw, 0.95, label=f"STORY {n}", people=ppl)
        text(s, ix, y + 1.06, iw, 0.36, P(en, size=18, bold=True, color=NAVY, space_after=0))
        text(s, ix, y + 1.42, iw, 0.32, P(ar, size=14, color=TEAL, bold=True, rtl=True, align="l", space_after=0))
        text(s, ix, y + 1.84, iw, 0.24, P("THE SITUATION", size=10, color=MUTED, bold=True, space_after=0))
        text(s, ix, y + 2.07, iw, 0.7, P(situation, size=13, space_after=0))
        text(s, ix, y + 2.8, iw, 0.24, P("THE OFFICER TYPES", size=10, color=MUTED, bold=True, space_after=0))
        box(
            s, ix, y + 3.04, iw, 0.78, fill=WHITE, line=TEAL, line_w=1.5, radius=0.12,
            paras=[
                P(q, size=15, bold=True, color=NAVY, rtl=True, space_after=1),
                P(gloss, size=10.5, color=MUTED, italic=True, align="r", rtl=False, space_after=0),
            ],
            anchor="middle", pad=0.12,
        )
        text(s, ix, y + 3.94, iw, 0.24, P("WHAT THEY GET", size=10, color=TEAL, bold=True, space_after=0))
        text(s, ix, y + 4.17, iw, 0.72, P(result, size=13, space_after=0))


def slide_safeguards(deck: Deck):
    s = deck.new(
        kicker="Safeguards",
        arabic="الضمانات والخصوصية",
        title="Built in from day one: appearance only, people decide, every search on record",
        notes="""
Report section 5, "Legal, privacy and ethics". This is not legal advice: a Jordanian lawyer should review before any pilot with real footage.
No face recognition: decision D5.
Lead, not proof: false matches are expected; the Robert Williams wrongful arrest and NIST's 10-100x demographic differentials in face recognition are the reasons for the design rule. Every result shows its confidence and why it matched, and is confirmed by people, two for any action.
Audit: role-based access, a case number per search, an append-only audit log; the index is deleted with the footage; faces blurred in exports; a data protection impact assessment before any real-footage pilot.
Data in Jordan: the PoC uses local models only, no hosted models even for testing (decision D4).
Jordan PDPL No. 24 of 2023: in force since 17 Mar 2024, full compliance since 17 Mar 2025. CCTV is personal data; biometric data is sensitive; consent must be explicit and documented, or an Art. 6 exception must apply; transfers abroad are restricted (Art. 15); breaches notified within 72 hours. Official translation: https://www.modee.gov.jo/EBV4.0/Root_Storage/EN/1/PDP_Law_-_English_Version-_officail_translation.pdf
""",
    )
    cards = [
        ("No face recognition", "لا للتعرف على الوجوه", "Searches clothing, colours, bags and size. It never identifies who someone is."),
        ("A lead, not proof", "نتيجة للتحقق لا دليل", "Each result shows its confidence and why it matched. A person confirms; two people before any action."),
        ("Every search on record", "سجل تدقيق", "A case number for each search and an append-only log: who searched, the exact query, what they opened."),
        ("Data stays in Jordan", "البيانات داخل الأردن", "Models run locally. The PoC sends nothing to any cloud AI service."),
        ("Jordan's data law (PDPL)", "قانون حماية البيانات الشخصية", "Law No. 24 of 2023: CCTV is personal data; consent or a legal exception; transfers abroad restricted; breaches reported within 72 hours."),
    ]
    cw, gap, y, h = 2.29, 0.17, 1.95, 3.95
    for i, (en, ar, body) in enumerate(cards):
        x = 0.6 + i * (cw + gap)
        box(s, x, y, cw, h, fill=PANEL, radius=0.05)
        shape(s, MSO_SHAPE.RECTANGLE, x, y, cw, 0.08, fill=AMBER if i == 0 else TEAL)
        if i == 0:
            shape(s, MSO_SHAPE.NO_SYMBOL, x + 0.22, y + 0.3, 0.55, 0.55, fill=AMBER)
        else:
            circle_label(s, x + 0.22, y + 0.3, 0.55, str(i + 1), fill=TEAL, size=15)
        text(s, x + 0.22, y + 1.0, cw - 0.4, 0.62, P(en, size=16, bold=True, color=NAVY, space_after=0, line_spacing=0.95))
        text(s, x + 0.22, y + 1.62, cw - 0.4, 0.36, P(ar, size=11.5, color=TEAL, bold=True, rtl=True, align="l", space_after=0))
        text(s, x + 0.22, y + 2.05, cw - 0.4, 1.85, P(body, size=13, space_after=0))
    box(
        s, 0.6, 6.08, 12.13, 0.78, fill=AMBER_LIGHT, radius=0.08,
        paras=[P([R("Also: ", bold=True, color=AMBER),
                  R("the search index is deleted with the footage, other people's faces are blurred in exports, and an impact assessment comes before any pilot on real footage. "
                    "Not legal advice: a Jordanian lawyer reviews first.")], size=13, space_after=0)],
        anchor="middle", pad=0.2,
    )


def mock_result(slide, x, y, w, h, *, place, time, why, verdict, verdict_color, people, fill=WHITE):
    """One result card of the mock business view, laid out right to left like the page."""
    box(slide, x, y, w, h, fill=fill, line=LINE, radius=0.06)
    fw = h * 1.3
    frame(slide, x + w - fw - 0.08, y + 0.07, fw, h - 0.14, people=people)
    bw = 0.72
    for j, (label, color) in enumerate((("✓ تأكيد", TEAL), ("✗ رفض", AMBER))):
        box(slide, x + 0.1 + j * (bw + 0.06), y + h - 0.38, bw, 0.28, fill=WHITE, line=color, radius=0.3,
            paras=[P(label, size=9.5, bold=True, color=color, rtl=True, align="c", space_after=0)], anchor="middle", pad=0.02)
    text(
        slide, x + 0.1, y + 0.06, w - fw - 0.3, h - 0.45,
        [
            P([R(place + "  ·  ", color=INK), R("G331", bold=True, color=NAVY)], size=11, rtl=True, space_after=0),
            P([R(time, color=MUTED), R("   " + why, color=TEAL, bold=True)], size=10, rtl=True, space_after=0),
            P(verdict, size=10, bold=True, color=verdict_color, rtl=True, space_after=0),
        ],
    )


def slide_demo(deck: Deck):
    s = deck.new(
        kicker="The demo",
        arabic="العرض التجريبي",
        title="What you will see live: search in Arabic, check the clips, decide",
        notes="""
What the demo shows, as it runs today.
The business view opens in Arabic, right to left; one button switches the whole page to English and back.
The officer types a description, optionally ticks cameras and sets a time window, and the query model shows what it understood, as fields such as woman and grey or white jacket.
Each result plays the clip at its moment, with camera, place and time, and shows why it matched, detail by detail.
A few seconds later the vision model double-checks the top results against each detail of the query: confirmed ones move up, rejected ones move down, unsure ones stay put.
The officer confirms or rejects each result; a result is a lead, not proof.
Every upload, search, confirm and reject goes into an audit log that cannot be edited or deleted.
A separate technical view shows the internals, step by step, for technical questions.
Where it runs (D3): one local computer, with no containers. The app and the AI models all run as native processes, started by plain scripts, so the models can use the computer's GPU.
Footage (D2, D7): recordings from public datasets, no live cameras: MEVA (Kitware and IARPA, CC BY 4.0, https://mevadata.org/resources/MEVA-data-license.txt), NVIDIA PhysicalAI-SmartSpaces synthetic store scenes (CC BY 4.0), and CAVIAR left-bag clips (Creative Commons BY-SA). The lost-child story uses synthetic people, so no real children appear.
The screen on the right is an illustration of the business view, not a screenshot.
""",
    )
    steps = [
        ("Arabic first, English on one click", "The page opens in Arabic. One button switches all of it to English."),
        ("Clips with camera, place and time", "Each result plays at its moment and shows why it matched, detail by detail."),
        ("A second look by a vision model", "Seconds later it double-checks the top results: confirmed ones move up, rejected ones down."),
        ("People decide, and every search is logged", "The officer confirms or rejects each result. Every search and decision goes into a log that cannot be edited."),
    ]
    numbered_steps(s, 0.6, 1.95, 5.3, steps, heights=[0.78, 0.78, 0.78, 0.98], gap=0.2, size=14)

    # Mock of the business view, right to left as on the page.
    mx, my, mw, mh = 6.3, 1.9, 6.43, 4.15
    box(s, mx, my, mw, mh, fill="EEF1F4", line=LINE, radius=0.03)
    shape(s, MSO_SHAPE.RECTANGLE, mx, my, mw, 0.36, fill=NAVY)
    text(s, mx + 1.3, my + 0.04, mw - 1.45, 0.28, P("البحث الذكي في تسجيلات المراقبة", size=12, color=WHITE, bold=True, rtl=True, space_after=0))
    box(s, mx + 0.12, my + 0.06, 0.95, 0.24, fill="153A5C", line="7FD1CC", radius=0.5,
        paras=[P("English", size=9.5, color=WHITE, bold=True, align="c", rtl=False, space_after=0)], anchor="middle", pad=0.02)
    text(s, mx + 0.2, my + 0.44, mw - 0.4, 0.26,
         P([R("البحث", bold=True, color=TEAL), R("     التسجيلات     سجل التدقيق", color=MUTED)], size=11, rtl=True, space_after=0))
    box(s, mx + 1.2, my + 0.76, mw - 1.4, 0.48, fill=WHITE, line=TEAL, line_w=1.25, radius=0.2,
        paras=[P("امراة تلبس جاكيت سكني او ابيض", size=15, bold=True, color=NAVY, rtl=True, space_after=0)], anchor="middle", pad=0.12)
    box(s, mx + 0.2, my + 0.76, 0.9, 0.48, fill=TEAL, radius=0.2,
        paras=[P("ابحث", size=14, bold=True, color=WHITE, align="c", rtl=True, space_after=0)], anchor="middle", pad=0.04)
    text(s, mx + 3.75, my + 1.36, mw - 3.95, 0.3, P("فهمنا من الوصف:", size=10.5, color=MUTED, bold=True, rtl=True, space_after=0))
    chips = [("امرأة", 0.75), ("سترة رمادي أو أبيض", 1.55)]
    cx = mx + 3.65
    for label, cwid in chips:
        cx -= cwid + 0.08
        box(s, cx, my + 1.34, cwid, 0.3, fill=TEAL_LIGHT, radius=0.5,
            paras=[P(label, size=10, color=NAVY, align="c", rtl=True, space_after=0)], anchor="middle", pad=0.02)
    rows = [
        dict(place="صالة انتظار محطة الحافلات", time="17:00:41", why="✓ امرأة  ✓ سترة", verdict="✓ أكّد النموذج البصري",
             verdict_color=TEAL, people=[(0.35, 0.82, dict(top="E8E1D3"))]),
        dict(place="صالة انتظار محطة الحافلات", time="17:01:16", why="✓ امرأة  ✓ سترة", verdict="✓ أكّد النموذج البصري",
             verdict_color=TEAL, people=[(0.2, 0.6, dict(top=GREY_TOP)), (0.55, 0.8, dict(top="C7CCD1", bag="C8A97E"))]),
        dict(place="صالة انتظار محطة الحافلات", time="17:03:05", why="≈ امرأة  ✓ سترة", verdict="✗ رفض النموذج البصري: نزلت إلى الأسفل",
             verdict_color=AMBER, people=[(0.4, 0.8, dict(top=GREY_TOP, bottom=BAG_BLACK))], fill=AMBER_LIGHT),
    ]
    for i, r in enumerate(rows):
        mock_result(s, mx + 0.2, my + 1.76 + i * 0.78, mw - 0.4, 0.7, **r)
    text(s, mx, my + mh + 0.04, mw, 0.26, P("Illustration of the demo's business view, not a screenshot.", size=10, color=MUTED, italic=True, space_after=0))

    box(
        s, 0.6, 6.3, 12.13, 0.62, fill=TEAL_LIGHT, radius=0.08,
        paras=[P([R("Runs on one computer: ", bold=True, color=NAVY),
                  R("the app and the AI models run as native processes, with no containers and no cloud AI. "
                    "Recorded public footage only: actors and synthetic people, no live cameras.")], size=13, space_after=0)],
        anchor="middle", pad=0.2,
    )


# ---------------------------------------------------------------- build


def set_theme_fonts(prs: Presentation):
    """Make Aptos and Dubai the theme fonts, so any default text matches the deck."""
    master = prs.slide_masters[0]
    for rel in master.part.rels.values():
        if rel.reltype.endswith("/theme"):
            part = rel.target_part
            root = etree.fromstring(part.blob)
            for group in ("a:majorFont", "a:minorFont"):
                g = root.find(f".//{qn(group)}")
                g.find(qn("a:latin")).set("typeface", LATIN_FONT)
                g.find(qn("a:cs")).set("typeface", ARABIC_FONT)
            part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


# The cover is not numbered; the five slides after it are.
SLIDES = [
    slide_title,
    slide_problem,
    slide_what,
    slide_stories,
    slide_safeguards,
    slide_demo,
]


def build(out: Path = OUT) -> Path:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    set_theme_fonts(prs)
    props = prs.core_properties
    props.title = "Arabic AI Video Search: PoC use cases"
    props.subject = "How security teams would use Arabic video search"
    props.author = ""
    props.last_modified_by = ""
    props.created = props.modified = dt.datetime(2026, 10, 17)
    props.revision = 1
    deck = Deck(prs, total=len(SLIDES) - 1)
    for make in SLIDES:
        make(deck)
    prs.save(out)
    return out


if __name__ == "__main__":
    print(build())
